# -*- coding: utf-8 -*-
"""Convierte el Excel del calendario escolar en el events.json que consume el kiosk."""
import hashlib
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

MESES = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
    "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11,
    "DICIEMBRE": 12,
}
MES_RE = re.compile(r"^\s*(" + "|".join(MESES) + r")\s+(\d{4})\s*$", re.IGNORECASE)
DIAS = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO"]

# Texto que vive en la rejilla pero no es un evento (encabezados, leyendas, notas al pie)
RUIDO_RE = re.compile(
    r"^(\d+[ºª°]?\s*TRIMESTRE|LUNES|MARTES|MI[ÉE]RCOLES|JUEVES|VIERNES|S[ÁA]BADO|DOMINGO"
    r"|FESTIVOS?|VACACIONES|D[ÍI]AS|LEYENDA|OBSERVACIONES|Calendario Escolar.*)$",
    re.IGNORECASE)

# Hojas a procesar: nombre en el libro -> etiqueta mostrada en el kiosk
CURSO_INICIO = 2026

HOJAS = {
    "Calendario": "General",
    "ERASMUS + REDES": "Erasmus",
    "EMERGENCIAS OP-A1": "Emergencias",
    "CENTROS de EMERGENCIAS OP-A": "Emergencias centros",
}


def norm(s):
    if s is None:
        return ""
    s = str(s).replace("\n", " ")
    return re.sub(r"\s+", " ", s).strip()


def sin_tildes(s):
    for a, b in zip("ÁÉÍÓÚÜÑ", "AEIOUUN"):
        s = s.replace(a, b)
    return s


def build_grid(ws):
    """Valores de la hoja resolviendo celdas combinadas (el ancla se propaga)."""
    grid = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                grid[(cell.row, cell.column)] = cell.value
    for rng in ws.merged_cells.ranges:
        anchor = grid.get((rng.min_row, rng.min_col))
        if anchor is None:
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid.setdefault((r, c), anchor)
    return grid


def resolve_date(day, weekday, ref_year, ref_month):
    """El número de día puede pertenecer al mes anterior, al del bloque o al siguiente.
    Se elige el candidato cuyo día de la semana coincide con su columna."""
    candidates = []
    for delta in (-1, 0, 1):
        m = ref_month + delta
        y = ref_year
        if m < 1:
            m, y = 12, y - 1
        elif m > 12:
            m, y = 1, y + 1
        try:
            candidates.append(date(y, m, day))
        except ValueError:
            pass
    for cand in candidates:
        if cand.weekday() == weekday:
            return cand
    return None


def parse_sheet(ws, label):
    grid = build_grid(ws)
    events = []
    day_meta = {}

    month_blocks = []
    for (r, c), v in grid.items():
        m = MES_RE.match(norm(v)) if isinstance(v, str) else None
        if m:
            month_blocks.append((r, c, MESES[m.group(1).upper()], int(m.group(2))))
    month_blocks.sort()

    for r, c, month, year in month_blocks:
        header_row = r + 1
        first_header = sin_tildes(norm(grid.get((header_row, c), "")).upper())
        if first_header != "LUNES":
            continue

        # El bloque termina donde empieza el siguiente bloque de mes en la misma columna
        next_rows = [br for br, bc, _, _ in month_blocks if bc == c and br > r]
        end_row = min(next_rows) - 2 if next_rows else ws.max_row

        rows = list(range(header_row + 1, end_row + 1))
        day_rows = []
        for rr in rows:
            nums = [cc for cc in range(c, c + 7)
                    if isinstance(grid.get((rr, cc)), (int, float))]
            if len(nums) >= 5:
                day_rows.append(rr)
        if not day_rows:
            continue

        # Bajo la última semana solo hay filas de texto; el resto de la hoja
        # (tablas informativas al pie) queda fuera del bloque.
        gaps = [b - a for a, b in zip(day_rows, day_rows[1:])]
        end_row = min(end_row, day_rows[-1] + (max(gaps) - 1 if gaps else 3))

        for i, rr in enumerate(day_rows):
            next_day_row = day_rows[i + 1] if i + 1 < len(day_rows) else end_row + 1
            week = {}
            for offset in range(7):
                cc = c + offset
                val = grid.get((rr, cc))
                if not isinstance(val, (int, float)) or isinstance(val, bool):
                    continue
                d = resolve_date(int(val), offset, year, month)
                if d:
                    week[offset] = d

            for offset, d in week.items():
                in_month = (d.month == month and d.year == year)
                cell = ws.cell(row=rr, column=c + offset)
                fill = cell.fill
                rgb = None
                if fill is not None and fill.fgColor is not None:
                    rgb = fill.fgColor.rgb if isinstance(fill.fgColor.rgb, str) else None
                meta = day_meta.setdefault(d.isoformat(), {})
                if in_month:
                    meta["fill"] = rgb

            # Filas de texto bajo la fila de días
            raw = {}
            for rr2 in range(rr + 1, next_day_row):
                for offset in range(7):
                    val = grid.get((rr2, c + offset))
                    if isinstance(val, str):
                        t = norm(val)
                        if t and t != "-" and not RUIDO_RE.match(t):
                            raw.setdefault(offset, [])
                            if t not in raw[offset]:
                                raw[offset].append(t)

            for offset, titles in raw.items():
                if offset not in week:
                    continue
                for t in titles:
                    events.append({
                        "date": week[offset].isoformat(),
                        "title": t,
                        "source": label,
                    })

    return events, day_meta


def contiguo(prev, d):
    """Días consecutivos, o viernes -> lunes siguiente (el finde no rompe el rango)."""
    gap = (d - prev).days
    return gap == 1 or (gap == 3 and prev.weekday() == 4)


def merge_spans(events):
    """Une eventos idénticos en días consecutivos en un único rango."""
    by_key = {}
    for e in events:
        by_key.setdefault((e["title"], e["source"], e.get("kind", "evento")), []).append(e["date"])

    merged = []
    for (title, source, kind), dates in by_key.items():
        dates = sorted({date.fromisoformat(d) for d in dates})
        start = prev = dates[0]
        for d in dates[1:]:
            if contiguo(prev, d):
                prev = d
                continue
            merged.append({"start": start.isoformat(), "end": prev.isoformat(),
                           "title": title, "source": source, "kind": kind})
            start = prev = d
        merged.append({"start": start.isoformat(), "end": prev.isoformat(),
                       "title": title, "source": source, "kind": kind})
    merged.sort(key=lambda e: (e["start"], e["title"]))
    for e in merged:
        raw = f"{e['source']}|{e['title']}|{e['start']}|{e['end']}"
        e["id"] = "e_" + hashlib.md5(raw.encode("utf-8")).hexdigest()[:10]
    return merged


def dedupe(events):
    """La misma fecha+título aparece en varias hojas; se conserva una sola vez."""
    prioridad = {label: i for i, label in enumerate(HOJAS.values())}
    mejor = {}
    for e in events:
        key = (e["date"], e["title"].upper())
        actual = mejor.get(key)
        if actual is None or prioridad.get(e["source"], 99) < prioridad.get(actual["source"], 99):
            mejor[key] = e
    return list(mejor.values())


def parse_legend(ws, curso_inicio):
    """Festivos (fecha suelta) y vacaciones (rango escrito en texto) de la leyenda."""
    grid = build_grid(ws)
    festivos = {}
    vacaciones = []

    festivos_hdr = [(r, c) for (r, c), v in grid.items()
                    if isinstance(v, str) and norm(v).upper() == "FESTIVOS"]
    for r, c in festivos_hdr:
        for rr in range(r + 2, r + 20):
            fecha = grid.get((rr, c))
            nombre = grid.get((rr, c + 1))
            if not hasattr(fecha, "date") or not isinstance(nombre, str):
                continue
            d = fecha.date() if hasattr(fecha, "date") else fecha
            # El libro trae algún año mal tecleado: se reancla al curso correcto
            year = curso_inicio if d.month >= 9 else curso_inicio + 1
            d = date(year, d.month, d.day)
            festivos[d.isoformat()] = norm(nombre).title()

    vac_hdr = [(r, c) for (r, c), v in grid.items()
               if isinstance(v, str) and norm(v).upper() == "VACACIONES"]
    for r, c in vac_hdr:
        for rr in range(r + 2, r + 20):
            rango = grid.get((rr, c))
            nombre = grid.get((rr, c + 1))
            if not isinstance(rango, str) or not isinstance(nombre, str):
                continue
            m = re.match(r"(\d+)\s+([A-ZÁÉÍÓÚÑ]+)\s*-\s*(\d+)\s+([A-ZÁÉÍÓÚÑ]+)",
                         norm(rango).upper())
            if not m:
                continue
            d1, m1, d2, m2 = int(m.group(1)), MESES.get(m.group(2)), int(m.group(3)), MESES.get(m.group(4))
            if not m1 or not m2:
                continue
            y1 = curso_inicio if m1 >= 9 else curso_inicio + 1
            y2 = curso_inicio if m2 >= 9 else curso_inicio + 1
            vacaciones.append({"start": date(y1, m1, d1).isoformat(),
                               "end": date(y2, m2, d2).isoformat(),
                               "title": norm(nombre).title()})
    vacaciones.sort(key=lambda v: v["start"])
    return festivos, vacaciones


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(
        r"C:\Users\aidan\Downloads\Calendario 2026-2027.xlsx")
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).parent.parent / "web" / "data" / "events.json"

    wb = openpyxl.load_workbook(src, data_only=True)
    all_events = []
    for sheet_name, label in HOJAS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ev, _ = parse_sheet(wb[sheet_name], label)
        all_events.extend(ev)

    all_events = dedupe(all_events)
    festivos, vacaciones = parse_legend(wb[list(HOJAS)[0]], CURSO_INICIO)

    def palabras(t):
        return {w for w in sin_tildes(t.upper()).split() if len(w) > 3}

    # El nombre del festivo ya se pinta desde `festivos`; no se repite como evento
    def es_festivo_repetido(e):
        f = festivos.get(e["date"])
        return f is not None and len(palabras(f) & palabras(e["title"])) >= 2

    all_events = [e for e in all_events if not es_festivo_repetido(e)]
    merged = merge_spans(all_events)

    out.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "curso": f"{CURSO_INICIO}-{CURSO_INICIO + 1}",
        "generated_from": src.name,
        "sources": list(HOJAS.values()),
        "festivos": festivos,
        "vacaciones": vacaciones,
        "events": merged,
    }
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(merged)} eventos, {len(festivos)} festivos, {len(vacaciones)} periodos de vacaciones -> {out}")


if __name__ == "__main__":
    main()
